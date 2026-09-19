import io
import json
import re
from datetime import date, datetime, timedelta
from pathlib import Path
from urllib.parse import parse_qs, urlencode, urljoin, urlparse

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook


BOARD_URL = "https://www.gwangju.ac.kr/bbs/?b_id=gwangju_jinwol_rm&mn=553&site=gwangju"
BASE_URL = "https://www.gwangju.ac.kr"
HEADERS = {
    "User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 "
                  "(KHTML, like Gecko) Chrome/153.0.0.0 Safari/537.36"
}
WEEKDAY_KO = ["월", "화", "수", "목", "금", "토", "일"]


def clean(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    text = str(value).replace("\r\n", "\n").replace("\r", "\n")
    # 셀 안 줄바꿈은 유지하고, 각 줄의 불필요한 공백만 정리합니다.
    return "\n".join(" ".join(line.split()) for line in text.split("\n")).strip()


def canonical_post_url(bs_idx):
    query = urlencode({
        "b_id": "gwangju_jinwol_rm",
        "mn": "553",
        "site": "gwangju",
        "type": "view",
        "bs_idx": str(bs_idx),
    })
    return f"{BASE_URL}/bbs/?{query}"


def get_latest_post():
    """식당메뉴 목록에서 가장 큰 bs_idx의 게시글을 최신 글로 선택합니다.

    제목 문구에만 의존하지 않아 학교가 제목 표현을 조금 바꿔도 동작하도록 했습니다.
    """
    r = requests.get(BOARD_URL, headers=HEADERS, timeout=15)
    r.raise_for_status()
    soup = BeautifulSoup(r.text, "html.parser")

    candidates = {}
    for a in soup.find_all("a", href=True):
        href = urljoin(BASE_URL, a.get("href", ""))
        parsed = urlparse(href)
        qs = parse_qs(parsed.query)
        bs_values = qs.get("bs_idx")
        if not bs_values:
            continue

        try:
            bs_idx = int(bs_values[0])
        except (TypeError, ValueError):
            continue

        # 식당메뉴 게시판이 아닌 링크는 제외합니다.
        b_id = (qs.get("b_id") or [""])[0]
        if b_id and b_id != "gwangju_jinwol_rm":
            continue

        own_text = " ".join(a.stripped_strings).strip()
        parent_text = " ".join(a.parent.stripped_strings).strip() if a.parent else ""
        row = a.find_parent("tr")
        row_text = " ".join(row.stripped_strings).strip() if row else ""
        title_text = own_text or parent_text or row_text

        # 동일 게시글 링크가 여러 개면 더 설명적인 텍스트를 보존합니다.
        current = candidates.get(bs_idx, "")
        if len(title_text) > len(current):
            candidates[bs_idx] = title_text

    if not candidates:
        raise RuntimeError("식단 게시글 링크를 찾지 못했습니다. 학교 홈페이지 구조가 바뀌었을 수 있습니다.")

    latest_idx = max(candidates)
    post_url = canonical_post_url(latest_idx)

    # 실제 상세 페이지에서 제목을 다시 확인합니다.
    r = requests.get(post_url, headers=HEADERS, timeout=15)
    r.raise_for_status()
    detail = BeautifulSoup(r.text, "html.parser")

    title = ""
    for selector in ["h3", "h4", ".bbs_title", ".view_title", ".subject", "title"]:
        for node in detail.select(selector):
            text = " ".join(node.stripped_strings).strip()
            if "메뉴" in text and ("학생" in text or "교직원" in text):
                title = text
                break
        if title:
            break

    if not title:
        page_text = "\n".join(detail.stripped_strings)
        match = re.search(
            r"(20\d{2}[.\-/]\s*\d{1,2}[.\-/]\s*\d{1,2}[^\n]{0,80}(?:학생정식|교직원)[^\n]{0,80}메뉴[^\n]*)",
            page_text,
        )
        if match:
            title = " ".join(match.group(1).split())

    if not title:
        title = candidates[latest_idx] or f"식당메뉴 게시글 #{latest_idx}"

    return {"title": title, "url": post_url, "bs_idx": latest_idx}


def find_excel_attachment(post_url):
    """상세 글에서 엑셀 첨부 다운로드 링크를 찾습니다."""
    r = requests.get(post_url, headers=HEADERS, timeout=15)
    r.raise_for_status()
    soup = BeautifulSoup(r.text, "html.parser")

    scored = []
    for a in soup.find_all("a", href=True):
        text = " ".join(a.stripped_strings).strip()
        href = urljoin(BASE_URL, a["href"])
        low = f"{text} {href}".lower()
        score = 0
        if ".xlsx" in low:
            score += 100
        elif ".xls" in low:
            score += 90
        if "type=download" in low:
            score += 50
        if "bf_idx=" in low:
            score += 40
        if "download" in low or "file_down" in low or "filedown" in low or "attach" in low:
            score += 20
        if score:
            scored.append((score, href, text))

    if not scored:
        raise RuntimeError("식단 엑셀 첨부파일 링크를 찾지 못했습니다.")

    scored.sort(key=lambda x: x[0], reverse=True)
    return scored[0][1]


def download_workbook(url):
    r = requests.get(url, headers=HEADERS, timeout=20)
    r.raise_for_status()
    if len(r.content) < 100:
        raise RuntimeError("첨부 식단표를 내려받았지만 파일 내용이 비어 있습니다.")
    return r.content


def workbook_matrix(binary):
    wb = load_workbook(io.BytesIO(binary), data_only=True)
    ws = wb[wb.sheetnames[0]]

    # 병합 셀의 값을 병합 범위 전체에 복사해 파싱을 안정화합니다.
    merged_values = {}
    for rng in ws.merged_cells.ranges:
        value = clean(ws.cell(rng.min_row, rng.min_col).value)
        if not value:
            continue
        for row in range(rng.min_row, rng.max_row + 1):
            for col in range(rng.min_col, rng.max_col + 1):
                merged_values[(row, col)] = value

    matrix = []
    for r in range(1, ws.max_row + 1):
        row = []
        for c in range(1, ws.max_column + 1):
            value = merged_values.get((r, c), clean(ws.cell(r, c).value))
            row.append(value)
        matrix.append(row)
    return matrix


def parse_title_dates(title):
    """게시글 제목의 2026.09.21~09.25 같은 범위에서 실제 날짜 목록을 만듭니다."""
    normalized = title.replace(" ", "")
    pattern = re.compile(
        r"(?P<y>20\d{2})[.\-/](?P<m>\d{1,2})[.\-/](?P<d>\d{1,2})"
        r"(?:\([^)]*\))?~"
        r"(?:(?P<y2>20\d{2})[.\-/])?(?P<m2>\d{1,2})[.\-/](?P<d2>\d{1,2})"
    )
    match = pattern.search(normalized)
    if not match:
        return []

    start = date(int(match.group("y")), int(match.group("m")), int(match.group("d")))
    end_year = int(match.group("y2") or match.group("y"))
    end = date(end_year, int(match.group("m2")), int(match.group("d2")))
    if end < start or (end - start).days > 14:
        return []

    dates = []
    current = start
    while current <= end:
        dates.append(current)
        current += timedelta(days=1)
    return dates


def value_matches_date(text, target):
    if not text:
        return False
    compact = re.sub(r"\s+", "", str(text))
    candidates = [
        target.strftime("%Y-%m-%d"),
        f"{target.year}.{target.month:02d}.{target.day:02d}",
        f"{target.year}.{target.month}.{target.day}",
        f"{target.month:02d}.{target.day:02d}",
        f"{target.month}.{target.day}",
        f"{target.month}/{target.day}",
        f"{target.month}월{target.day}일",
        f"{target.day}일",
    ]
    return any(c.replace(" ", "") in compact for c in candidates)


def is_noise(text):
    if not text:
        return True
    compact = re.sub(r"\s+", "", text).lower()
    exact_noise = {
        "lunch", "중식", "점심", "메뉴", "menu", "학생정식", "교직원", "교직원정식",
        "월", "화", "수", "목", "금", "토", "일", "mon", "tue", "wed", "thu", "fri",
    }
    if compact in exact_noise:
        return True
    if re.fullmatch(r"\d+(?:,\d{3})*원?", compact):
        return True
    if re.fullmatch(r"\d+(?:\.\d+)?k?cal", compact):
        return True
    if re.fullmatch(r"20\d{2}[-./]\d{1,2}[-./]\d{1,2}", compact):
        return True
    return False


def split_menu_text(text):
    if not text:
        return []
    # 메뉴가 한 셀 안에서 줄바꿈/쉼표/가운뎃점으로 묶여 있는 경우를 분리합니다.
    chunks = re.split(r"\n+|\s*[•·]\s*|\s*,\s*", text)
    result = []
    for chunk in chunks:
        value = " ".join(chunk.split()).strip(" -/|")
        if value and not is_noise(value):
            result.append(value)
    return result


def row_contains(matrix, row_index, keywords):
    if row_index < 0 or row_index >= len(matrix):
        return False
    text = " ".join(matrix[row_index]).replace(" ", "")
    return any(k.replace(" ", "") in text for k in keywords)


def normalized_row_text(row):
    return re.sub(r"\s+", "", " ".join(row)).lower()


def choose_horizontal_date_header(matrix, dates, anchor_keywords=None):
    """날짜/요일이 가로로 놓인 행을 찾습니다.

    anchor_keywords가 있으면 해당 메뉴 구역과 가까운 날짜 행을 우선합니다.
    """
    candidates = []
    normalized_keywords = [re.sub(r"\s+", "", k).lower() for k in (anchor_keywords or [])]

    for r, row in enumerate(matrix):
        mapping = {}
        for c, value in enumerate(row):
            for d in dates:
                if value_matches_date(value, d):
                    mapping[c] = d
                    break
            compact = re.sub(r"\s+", "", value)
            if c not in mapping:
                for d in dates:
                    w = WEEKDAY_KO[d.weekday()]
                    if compact in {w, f"{w}요일", f"({w})"}:
                        mapping[c] = d
                        break

        distinct = len(set(mapping.values()))
        if distinct < 2:
            continue

        proximity = 0
        if normalized_keywords:
            for rr in range(max(0, r - 8), min(len(matrix), r + 12)):
                row_text = normalized_row_text(matrix[rr])
                if any(k in row_text for k in normalized_keywords):
                    proximity += max(1, 12 - abs(rr - r))
        else:
            for rr in range(max(0, r - 5), min(len(matrix), r + 6)):
                if row_contains(matrix, rr, ["학생정식", "학생 식당", "학생"]):
                    proximity += max(1, 6 - abs(rr - r))

        candidates.append((distinct, proximity, r, mapping))

    if not candidates:
        return None
    candidates.sort(key=lambda x: (x[1], x[0]), reverse=True)
    _, _, row_index, mapping = candidates[0]
    return row_index, mapping


def infer_day_columns(matrix, dates):
    """날짜 헤더를 찾지 못했을 때 식단 데이터가 많은 연속 열을 날짜에 대응합니다."""
    if not matrix or not dates:
        return {}
    max_cols = max(len(row) for row in matrix)
    scores = []
    for c in range(max_cols):
        count = 0
        for row in matrix:
            if c < len(row):
                value = row[c]
                if value and not is_noise(value):
                    count += 1
        scores.append((count, c))

    useful = [c for count, c in scores if count >= 2]
    if len(useful) < len(dates):
        return {}

    best = None
    for start_pos in range(0, len(useful) - len(dates) + 1):
        cols = useful[start_pos:start_pos + len(dates)]
        continuity = sum(1 for a, b in zip(cols, cols[1:]) if b == a + 1)
        score_by_col = {c: count for count, c in scores}
        density = sum(score_by_col.get(c, 0) for c in cols)
        candidate = (continuity, density, cols)
        if best is None or candidate[:2] > best[:2]:
            best = candidate
    if not best:
        return {}
    return {c: d for c, d in zip(best[2], dates)}


SECTION_ALIASES = {
    "student": ["학생정식", "학생 정식", "학생식당", "학생 식당"],
    "food": ["푸드단품", "푸드 단품", "단품메뉴", "단품 메뉴", "푸드코트", "푸드 코트"],
    "faculty": ["교직원정식", "교직원 정식", "교직원"],
}


def find_section_anchors(matrix, aliases):
    normalized = [re.sub(r"\s+", "", a).lower() for a in aliases]
    found = []
    for r, row in enumerate(matrix):
        text = normalized_row_text(row)
        if any(alias in text for alias in normalized):
            found.append(r)
    return found


def section_bounds(matrix, header_row, aliases, all_section_aliases=None):
    """특정 메뉴 구역의 실제 행 범위를 찾습니다."""
    anchors = find_section_anchors(matrix, aliases)
    if not anchors:
        return None

    # 날짜 행에 가장 가까우면서 보통 그 아래에 있는 라벨을 우선합니다.
    below = [r for r in anchors if r >= header_row - 1]
    anchor = min(below, key=lambda r: abs(r - header_row)) if below else min(anchors, key=lambda r: abs(r - header_row))
    start = anchor + 1

    stop_aliases = all_section_aliases or []
    stop_norm = [re.sub(r"\s+", "", a).lower() for a in stop_aliases]
    end = min(len(matrix), start + 14)
    for r in range(start, min(len(matrix), start + 20)):
        text = normalized_row_text(matrix[r])
        if r > start and any(alias in text for alias in stop_norm):
            end = r
            break
        # 새 Lunch 블록이 시작되면 현재 메뉴 구역 종료로 봅니다.
        if r > start + 1 and text in {"lunch", "중식", "점심"}:
            end = r
            break
    return start, end


def collect_section_by_columns(matrix, dates, header_row, col_to_date, aliases):
    all_aliases = []
    for values in SECTION_ALIASES.values():
        all_aliases.extend(values)

    bounds = section_bounds(matrix, header_row, aliases, all_aliases)
    if not bounds:
        return {d.isoformat(): [] for d in dates}

    start, end = bounds
    items = {d.isoformat(): [] for d in dates}
    seen = {d.isoformat(): set() for d in dates}

    for r in range(start, end):
        row = matrix[r]
        row_text = normalized_row_text(row)
        # 다른 섹션 라벨 행은 데이터로 포함하지 않습니다.
        if any(re.sub(r"\s+", "", a).lower() in row_text for a in all_aliases):
            continue

        for col, d in col_to_date.items():
            if col >= len(row):
                continue
            for item in split_menu_text(row[col]):
                compact = re.sub(r"\s+", "", item)
                if any(value_matches_date(item, x) for x in dates):
                    continue
                if any(token in compact for token in [
                    "학생정식", "교직원", "푸드단품", "단품메뉴", "운영시간", "가격", "원산지"
                ]):
                    continue
                key = d.isoformat()
                if item not in seen[key]:
                    items[key].append(item)
                    seen[key].add(item)
    return items


def collect_vertical_section(matrix, dates, aliases):
    """날짜가 행 방향인 표에서 특정 메뉴 구역만 읽는 보조 파서."""
    result = {d.isoformat(): [] for d in dates}
    anchors = find_section_anchors(matrix, aliases)
    if not anchors:
        return result

    all_aliases = []
    for values in SECTION_ALIASES.values():
        all_aliases.extend(values)
    all_norm = [re.sub(r"\s+", "", a).lower() for a in all_aliases]

    for anchor in anchors:
        end = min(len(matrix), anchor + 18)
        for rr in range(anchor + 1, end):
            row_text = normalized_row_text(matrix[rr])
            if rr > anchor + 1 and any(a in row_text for a in all_norm):
                break
            target = None
            for value in matrix[rr]:
                for d in dates:
                    if value_matches_date(value, d):
                        target = d
                        break
                if target:
                    break
            if not target:
                continue

            key = target.isoformat()
            for value in matrix[rr]:
                for item in split_menu_text(value):
                    if value_matches_date(item, target) or is_noise(item):
                        continue
                    if item not in result[key]:
                        result[key].append(item)
    return result


def parse_meal_categories(matrix, title):
    dates = parse_title_dates(title)
    if not dates:
        return [], None, "게시글 제목에서 날짜 범위를 읽지 못했습니다."

    # 학생정식과 푸드단품은 표 안에서 위치가 다를 수 있어 날짜 헤더를 각각 찾습니다.
    student_header = choose_horizontal_date_header(matrix, dates, SECTION_ALIASES["student"])
    food_header = choose_horizontal_date_header(matrix, dates, SECTION_ALIASES["food"])

    if student_header:
        student_header_row, student_cols = student_header
    else:
        student_header_row, student_cols = 0, infer_day_columns(matrix, dates)

    if food_header:
        food_header_row, food_cols = food_header
    else:
        # 같은 주간표 안에 있으면 학생정식과 같은 날짜 열을 쓰는 경우가 많아 우선 재사용합니다.
        food_header_row, food_cols = student_header_row, dict(student_cols)

    student = collect_section_by_columns(
        matrix, dates, student_header_row, student_cols, SECTION_ALIASES["student"]
    ) if student_cols else {d.isoformat(): [] for d in dates}

    food = collect_section_by_columns(
        matrix, dates, food_header_row, food_cols, SECTION_ALIASES["food"]
    ) if food_cols else {d.isoformat(): [] for d in dates}

    # 가로 파싱에서 못 찾은 카테고리만 세로형 파서를 보조적으로 사용합니다.
    if not any(student.values()):
        student = collect_vertical_section(matrix, dates, SECTION_ALIASES["student"])
    if not any(food.values()):
        food = collect_vertical_section(matrix, dates, SECTION_ALIASES["food"])

    days = []
    for d in dates:
        key = d.isoformat()
        days.append({
            "date": key,
            "weekday": WEEKDAY_KO[d.weekday()],
            "student_items": student.get(key, []),
            "food_items": food.get(key, []),
            # 이전 화면 코드와의 호환성을 위해 유지합니다.
            "items": student.get(key, []),
        })

    today = date.today()
    available = [
        d for d in dates
        if student.get(d.isoformat()) or food.get(d.isoformat())
    ]
    if today in available:
        default_date = today
    else:
        future = [d for d in available if d >= today]
        default_date = future[0] if future else (available[-1] if available else dates[0])

    found_food = any(food.values())
    if found_food:
        parse_note = "광주대학교 공식 식단표에서 학생정식과 푸드단품을 구분해 불러왔습니다."
    else:
        parse_note = "학생정식은 확인했지만 이번 식단표에서 푸드단품 구역을 찾지 못했습니다."

    return days, default_date.isoformat(), parse_note

def rows_to_preview(matrix, limit=25):
    preview = []
    for idx, row in enumerate(matrix[:limit], start=1):
        cells = list(row)
        while cells and not cells[-1]:
            cells.pop()
        if any(cells):
            preview.append({"row": idx, "cells": cells})
    return preview

OUTPUT_PATH = Path(__file__).with_name("menu.json")


def build_menu_payload():
    post = get_latest_post()
    attachment = find_excel_attachment(post["url"])
    binary = download_workbook(attachment)
    matrix = workbook_matrix(binary)
    days, default_date, parse_note = parse_meal_categories(matrix, post["title"])

    return {
        "ok": True,
        "source": BOARD_URL,
        "post": post,
        "days": days,
        "default_date": default_date,
        "parse_note": parse_note,
        "last_updated": datetime.now().astimezone().isoformat(timespec="seconds"),
    }


def main():
    try:
        payload = build_menu_payload()
        tmp_path = OUTPUT_PATH.with_suffix(".json.tmp")
        tmp_path.write_text(
            json.dumps(payload, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        tmp_path.replace(OUTPUT_PATH)

        print("")
        print("✅ 광주대학교 식단 업데이트 완료")
        print(f"저장 파일: {OUTPUT_PATH}")
        print(f"게시글: {payload['post']['title']}")
        print("")
        print("이제 GitHub에서 menu.json 파일 하나만 업로드하고 Commit changes를 누르면 됩니다.")
        return 0
    except Exception as exc:
        print("")
        print("❌ 식단 업데이트에 실패했습니다.")
        print(str(exc))
        print("")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
